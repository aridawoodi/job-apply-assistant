"""Builds fixtures/casey-morgan-resume.docx/.pdf: a made-up person with 4 jobs and tricky dates
(Jan 2021 – Present, 2019–2021, the word "current" inside a bullet, overlapping jobs, start-year only)."""
import os, subprocess
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
HERE = os.path.dirname(os.path.abspath(__file__)); OUT = os.path.join(HERE, 'fixtures')
d = Document(); s = d.sections[0]
s.left_margin = s.right_margin = Inches(0.8); s.top_margin = s.bottom_margin = Inches(0.7)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)
def para(t, size=10.5, bold=False, center=False, after=2):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold; r.font.size = Pt(size)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after); return p
def heading(t): p = para(t, 11.5, True, after=3); p.paragraph_format.space_before = Pt(8)
def entry(left, right):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left); r.bold = True; p.add_run('\t' + right)
def bullet(t): p = d.add_paragraph(t, style='List Bullet'); p.paragraph_format.space_after = Pt(1)
para('Casey Morgan', 20, True, True, 0)
para('Lakeview, OR | (555) 010-0177 | casey.morgan@example.com', 9.5, center=True, after=6)
heading('EXPERIENCE')
entry('Senior Support Engineer, Bluefin Devices', 'Jan 2021 – Present')
bullet('Lead a team of 6 engineers handling escalations for 2 product lines.')
bullet('Built a triage dashboard used by 3 support teams.')
entry('Weekend Instructor, Lakeview Code Club', '2022 – Present')
bullet('Teach intro Python to adults on weekends.')
entry('Support Engineer, Harbor Cloud', '2019–2021')
bullet('Migrated the current ticket system to a new queue, cutting response time.')
bullet('Wrote runbooks for the on-call rotation.')
entry('IT Technician, Riverside Clinic', '2016')
bullet('Kept 40 workstations patched and backed up.')
heading('EDUCATION')
entry('B.S. Information Systems, Lakeview State University', '2012 – 2016')
heading('SKILLS')
para('Python, SQL, Linux, Zendesk, Jira')
os.makedirs(OUT, exist_ok=True)
docx = os.path.join(OUT, 'casey-morgan-resume.docx'); d.save(docx)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', OUT, docx], check=True, capture_output=True)
print('wrote', docx)
